"""Отчёт «План / факт»: проект → задача → сотрудник и сотрудник → проект → задача.

План — пользовательское поле задачи Битрикс24 («Оценка», часы), факт —
списания приложения. Отчёт РАБОЧИЙ, а не слепок: оценки читаются с портала
при каждом формировании, поэтому правка оценки в карточке задачи видна сразу
после повторного «Сформировать».

Отчёт включён НЕ для всех порталов. Поле оценки у каждого портала своё
(код пользовательского поля генерирует Битрикс24), а сама договорённость
«план ведём в поле задачи» — внутренняя практика одного портала, не функция
продукта. Поэтому доступ определяется белым списком PLAN_FACT_PORTALS по
member_id; на остальных порталах ручки отвечают 404, а пункта меню нет.

Про удвоение плана. Оценку ставят и на этап, и на его подзадачи (этап 960 ч =
200 + 160 + 160 + 250 + 190 по исполнителям). Прямая сумма поля удвоила бы
план, поэтому у узла берётся СВОЯ оценка, а сумма подзадач — только когда
своей нет. Расхождение своей оценки с суммой подзадач отдаётся наружу
(children_plan_hours) — это сигнал «этап и подзадачи спланированы по-разному».

Про остаток. Часы списывают и в задачи без оценки. Сравнивать с планом весь
факт нельзя — получится «выполнение 260%». Поэтому рядом с фактом считается
planned_fact_hours: факт только по задачам, у которых есть план (свой или
унаследованный от этапа). Остаток = план − planned_fact_hours.
"""

import io
import logging
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Set

import openpyxl
from django.conf import settings
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .employee_ids import extract_bitrix_user_id, normalize_employee_id, resolve_employee_name
from .report_excel import ExportTooLargeError, MAX_EXPORT_ROWS
from .report_queries import resolve_current_group_for_row, resolve_task_titles_for_row

logger = logging.getLogger(__name__)

#: Порталы, где отчёт включён: member_id -> настройки.
#: Сейчас это только mainsoft.bitrix24.ru. Переопределяется настройкой Django
#: PLAN_FACT_PORTALS (тем же словарём) — для тестов и стендов.
DEFAULT_PLAN_FACT_PORTALS: Dict[str, Dict[str, str]] = {
    "ac90d2405edacb6e7e65074528ce1af8": {"estimate_field": "UF_TASKS_TASK_1757361570662"},
}

#: Ключ в ответе GET /api/features. Отдаётся ТОЛЬКО порталам из белого списка.
PLAN_FACT_FEATURE_KEY = "plan_fact_report"

MAX_PLAN_PAGES = 60          # 3000 задач с оценкой — заведомо выше реального
MAX_ANCESTOR_ROUNDS = 8      # глубина вложенности задач
MAX_CHILD_PAGES = 40         # страниц подзадач на один уровень
ID_CHUNK = 50

NO_PROJECT_KEY = "none"
UNKNOWN_TASK_ID = "unknown"


def plan_fact_settings(account: Any) -> Optional[Dict[str, str]]:
    """Настройки отчёта для портала аккаунта либо None, если отчёт выключен."""
    portals = getattr(settings, "PLAN_FACT_PORTALS", None)
    if portals is None:
        portals = DEFAULT_PLAN_FACT_PORTALS
    member_id = str(getattr(account, "member_id", "") or "").strip()
    if not member_id:
        return None
    config = portals.get(member_id)
    if not config or not config.get("estimate_field"):
        return None
    return config


def plan_fact_feature_payload() -> Dict[str, str]:
    """Запись для /api/features в том же формате, что у остальных функций."""
    return {"state": "on", "access": "full", "status": "active"}


def _camel_field(code: str) -> str:
    """UF_TASKS_TASK_123 -> ufTasksTask123: так tasks.task.list отдаёт ключи."""
    parts = [part for part in str(code).split("_") if part]
    if not parts:
        return ""
    return parts[0].lower() + "".join(part[:1].upper() + part[1:].lower() for part in parts[1:])


def _clean(value: Any) -> str:
    text = str(value if value is not None else "").strip()
    return "" if text in ("0", "None", "null") else text


def _to_hours(value: Any) -> float:
    try:
        number = float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return 0.0
    return number if number > 0 else 0.0


def _extract_tasks(response: Any) -> List[Dict[str, Any]]:
    if not isinstance(response, dict):
        return []
    result = response.get("result")
    if isinstance(result, dict):
        tasks = result.get("tasks")
        return tasks if isinstance(tasks, list) else []
    return result if isinstance(result, list) else []


def _normalize_task(raw: Mapping[str, Any], field: str) -> Optional[Dict[str, Any]]:
    task_id = _clean(raw.get("id") or raw.get("ID"))
    if not task_id:
        return None
    group = raw.get("group") if isinstance(raw.get("group"), dict) else {}
    responsible = raw.get("responsible") if isinstance(raw.get("responsible"), dict) else {}
    estimate = raw.get(_camel_field(field))
    if estimate is None:
        estimate = raw.get(field)
    return {
        "id": task_id,
        "title": str(raw.get("title") or raw.get("TITLE") or "").strip(),
        "parent_id": _clean(raw.get("parentId") or raw.get("PARENT_ID")),
        "group_id": _clean(raw.get("groupId") or raw.get("GROUP_ID")),
        "group_name": str(group.get("name") or "").strip(),
        "plan": _to_hours(estimate),
        "responsible_id": extract_bitrix_user_id(raw.get("responsibleId") or raw.get("RESPONSIBLE_ID")),
        "responsible_name": str(responsible.get("name") or "").strip(),
        "closed": str(raw.get("status") or raw.get("STATUS") or "") == "5",
        # Задача с оценкой или её потомок: состав таких поддеревьев известен
        # ПОЛНОСТЬЮ и актуально (см. fetch_plan_tasks).
        "in_plan_tree": False,
    }


def fetch_plan_tasks(client: Any, field: str) -> Dict[str, Dict[str, Any]]:
    """Задачи с заполненной оценкой и вся цепочка их родителей.

    Родители нужны, чтобы плановая задача без единого списания встала в дерево
    на своё место (под этап), а не всплыла в корень проекта.

    Потомки нужны для остатка. Списание помнит иерархию задачи СНИМКОМ на
    момент записи: подзадача, которую завели или перенесли под этап позже,
    по снимку остаётся вне этапа, и её часы не попали бы в «факт по плану».
    Поэтому поддеревья плановых задач читаются с портала целиком, и для них
    действует текущая иерархия, а не снимок (флаг in_plan_tree).

    Исключения наружу НЕ глушатся: вызывающий решает, что показать человеку.
    """
    select = ["ID", "TITLE", "GROUP_ID", "PARENT_ID", "RESPONSIBLE_ID", "STATUS", field]
    tasks: Dict[str, Dict[str, Any]] = {}

    start = 0
    for _ in range(MAX_PLAN_PAGES):
        response = client._bitrix_token.call_method(
            "tasks.task.list",
            {"filter": {f">{field}": 0}, "select": select, "order": {"ID": "asc"}, "start": start},
        )
        for raw in _extract_tasks(response):
            task = _normalize_task(raw, field)
            if task:
                tasks[task["id"]] = task
        next_value = response.get("next") if isinstance(response, dict) else None
        try:
            next_start = int(next_value)
        except (TypeError, ValueError):
            break
        if next_start <= start:
            break
        start = next_start
    else:
        logger.warning("plan_fact: достигнут предел страниц (%s) при чтении оценок", MAX_PLAN_PAGES)

    for _ in range(MAX_ANCESTOR_ROUNDS):
        missing = sorted({t["parent_id"] for t in tasks.values() if t["parent_id"]} - set(tasks))
        if not missing:
            break
        fetched = 0
        for offset in range(0, len(missing), ID_CHUNK):
            response = client._bitrix_token.call_method(
                "tasks.task.list",
                {"filter": {"ID": missing[offset:offset + ID_CHUNK]}, "select": select},
            )
            for raw in _extract_tasks(response):
                task = _normalize_task(raw, field)
                if task and task["id"] not in tasks:
                    tasks[task["id"]] = task
                    fetched += 1
        if not fetched:  # родитель удалён или недоступен — дальше не ходим
            break

    frontier = sorted(task_id for task_id, task in tasks.items() if task["plan"] > 0)
    for task_id in frontier:
        tasks[task_id]["in_plan_tree"] = True
    for _ in range(MAX_ANCESTOR_ROUNDS):
        if not frontier:
            break
        found: List[str] = []
        for offset in range(0, len(frontier), ID_CHUNK):
            start = 0
            for _page in range(MAX_CHILD_PAGES):
                response = client._bitrix_token.call_method(
                    "tasks.task.list",
                    {
                        "filter": {"PARENT_ID": frontier[offset:offset + ID_CHUNK]},
                        "select": select, "order": {"ID": "asc"}, "start": start,
                    },
                )
                for raw in _extract_tasks(response):
                    task = _normalize_task(raw, field)
                    if not task:
                        continue
                    known = tasks.get(task["id"])
                    if known is None:
                        tasks[task["id"]] = known = task
                    if not known["in_plan_tree"]:
                        known["in_plan_tree"] = True
                        found.append(known["id"])
                next_value = response.get("next") if isinstance(response, dict) else None
                try:
                    next_start = int(next_value)
                except (TypeError, ValueError):
                    break
                if next_start <= start:
                    break
                start = next_start
        frontier = sorted(found)
    return tasks


def _employee_allowed(employee_id: str, selected: Set[str], mode: str) -> bool:
    if not selected:
        return True
    return (employee_id not in selected) if mode == "exclude" else (employee_id in selected)


def _project_allowed(group_id: str, selected: Set[str], mode: str, archived: Set[str]) -> bool:
    if group_id in archived:
        return False
    if not selected:
        return True
    return (group_id not in selected) if mode == "exclude" else (group_id in selected)


def build_plan_fact_report(
    rows: Iterable[Mapping[str, Any]],
    plan_tasks: Mapping[str, Mapping[str, Any]],
    *,
    user_map: Optional[Mapping[str, str]] = None,
    task_lookup: Optional[Mapping[str, Mapping[str, str]]] = None,
    project_name_by_group: Optional[Mapping[str, str]] = None,
    employee_ids: Sequence[str] = (),
    employee_mode: str = "include",
    project_group_ids: Sequence[str] = (),
    project_mode: str = "include",
    archived_group_ids: Sequence[str] = (),
) -> Dict[str, List[Dict[str, Any]]]:
    """Строки списаний + задачи с оценкой -> два дерева отчёта.

    rows уже отфильтрованы запросом (период, сотрудники, проекты). Фильтры
    передаются сюда повторно ради ПЛАНА: плановые задачи приходят с портала и
    запросом к списаниям не отсечены. Фильтр сотрудников к плану применяется
    по ответственному задачи.
    """
    user_map = user_map or {}
    names_by_group = dict(project_name_by_group or {})
    selected_employees = {extract_bitrix_user_id(v) for v in employee_ids if extract_bitrix_user_id(v)}
    selected_projects = {_clean(v) for v in project_group_ids if _clean(v)}
    archived = {_clean(v) for v in archived_group_ids if _clean(v)}

    nodes: Dict[str, Dict[str, Any]] = {}
    project_names: Dict[str, str] = {}

    def ensure(task_id: str, title: str, parent: str, project: str) -> Dict[str, Any]:
        node = nodes.get(task_id)
        if node is None:
            node = nodes[task_id] = {
                "id": task_id, "title": title, "parent": parent, "project": project,
                "own_plan": 0.0, "responsible_id": "", "responsible_name": "", "closed": False,
                "emps": {},
            }
        return node

    # --- факт ---
    for row in rows:
        hours = float(row.get("hours") or 0)
        if not hours:
            continue
        group_id = resolve_current_group_for_row(row, task_lookup) or _clean(row.get("project_id"))
        if group_id:
            project = group_id
            name = names_by_group.get(group_id) or str(row.get("project_title") or "").strip()
        else:
            name = str(row.get("project_title") or "").strip()
            project = f"name:{name}" if name else NO_PROJECT_KEY
        if name and project not in project_names:
            project_names[project] = name

        chain = [_clean(v) for v in (row.get("task_hierarchy_ids") or [])]
        titles = resolve_task_titles_for_row(row, task_lookup)
        if not any(chain):
            task_id = _clean(row.get("task_id"))
            chain = [task_id or UNKNOWN_TASK_ID]
            titles = titles[-1:] if task_id and titles else ["Без задачи"]

        parent = ""
        leaf = None
        for index, task_id in enumerate(chain):
            if not task_id:
                continue
            title = titles[index] if index < len(titles) and titles[index] else f"Задача {task_id}"
            key = task_id if task_id != UNKNOWN_TASK_ID else f"{UNKNOWN_TASK_ID}:{project}"
            leaf = ensure(key, title, parent, project)
            parent = key
        if leaf is None:
            continue
        employee = normalize_employee_id(row.get("employee_id"))
        bucket = leaf["emps"].setdefault(employee, [0.0, 0.0])
        bucket[0] += hours
        if row.get("is_billable"):
            bucket[1] += hours

    # --- план: актуальное состояние задачи с портала главнее снимка ---
    for task_id, task in plan_tasks.items():
        group_id = _clean(task.get("group_id"))
        if not _project_allowed(group_id, selected_projects, project_mode, archived):
            continue
        project = group_id or NO_PROJECT_KEY
        if project not in project_names:
            name = names_by_group.get(group_id) or str(task.get("group_name") or "").strip()
            if name:
                project_names[project] = name
        node = ensure(task_id, "", "", project)
        node["title"] = str(task.get("title") or "") or node["title"] or f"Задача {task_id}"
        node["parent"] = _clean(task.get("parent_id"))
        node["project"] = project
        node["responsible_id"] = str(task.get("responsible_id") or "")
        node["responsible_name"] = (
            user_map.get(node["responsible_id"]) or str(task.get("responsible_name") or "")
        )
        node["closed"] = bool(task.get("closed"))
        if _employee_allowed(node["responsible_id"], selected_employees, employee_mode):
            node["own_plan"] = float(task.get("plan") or 0)

    # --- связи: родитель в другом проекте или вне выборки — узел в корень ---
    # Состав поддеревьев плановых задач известен с портала полностью. Если
    # снимок списания помещает туда задачу, которой там сейчас нет, — её
    # перенесли, и к этапу она больше не относится.
    plan_tree_ids = {task_id for task_id, task in plan_tasks.items() if task.get("in_plan_tree")}
    children: Dict[str, List[str]] = {}
    roots: Dict[str, List[str]] = {}
    for task_id, node in nodes.items():
        if node["parent"] in plan_tree_ids and task_id not in plan_tasks:
            node["parent"] = ""
        parent = nodes.get(node["parent"]) if node["parent"] else None
        if parent is not None and parent["project"] == node["project"] and node["parent"] != task_id:
            children.setdefault(node["parent"], []).append(task_id)
        else:
            node["parent"] = ""
            roots.setdefault(node["project"], []).append(task_id)

    def title_key(task_id: str):
        return nodes[task_id]["title"].lower()

    def employee_rows(emps: Mapping[str, List[float]]) -> List[Dict[str, Any]]:
        result = [
            {
                "type": "employee", "id": emp_id, "name": resolve_employee_name(user_map, emp_id),
                "total_hours": round(v[0], 2), "billable_hours": round(v[1], 2),
                "non_billable_hours": round(v[0] - v[1], 2),
            }
            for emp_id, v in emps.items()
        ]
        result.sort(key=lambda item: -item["total_hours"])
        return result

    planned_parents: Set[str] = set()   # задачи, под которыми есть план глубже

    def calc(task_id: str, seen: frozenset) -> Optional[Dict[str, Any]]:
        if task_id in seen:  # цикл в данных — рвём
            return None
        node = nodes[task_id]
        kids = [k for k in (calc(c, seen | {task_id}) for c in sorted(children.get(task_id, []), key=title_key)) if k]
        own_total = sum(v[0] for v in node["emps"].values())
        own_billable = sum(v[1] for v in node["emps"].values())
        total = own_total + sum(k["total_hours"] for k in kids)
        billable = own_billable + sum(k["billable_hours"] for k in kids)
        kids_plan = sum(k["plan_hours"] for k in kids)
        plan = node["own_plan"] or kids_plan
        if kids_plan:
            planned_parents.add(task_id)
        if not total and not plan:
            return None
        return {
            "type": "task", "id": task_id, "name": node["title"],
            "plan_hours": round(plan, 2),
            "own_plan_hours": round(node["own_plan"], 2),
            "children_plan_hours": round(kids_plan, 2),
            "total_hours": round(total, 2), "billable_hours": round(billable, 2),
            "non_billable_hours": round(total - billable, 2),
            "planned_fact_hours": round(total if node["own_plan"] else sum(k["planned_fact_hours"] for k in kids), 2),
            "responsible_name": node["responsible_name"], "is_closed": node["closed"],
            "children": kids, "employees": employee_rows(node["emps"]),
        }

    projects: List[Dict[str, Any]] = []
    for project, task_ids in roots.items():
        kids = [k for k in (calc(t, frozenset()) for t in sorted(task_ids, key=title_key)) if k]
        if not kids:
            continue
        total = sum(k["total_hours"] for k in kids)
        billable = sum(k["billable_hours"] for k in kids)
        projects.append({
            "type": "project", "id": project,
            "name": project_names.get(project) or ("Без проекта" if project == NO_PROJECT_KEY else f"Проект {project}"),
            "plan_hours": round(sum(k["plan_hours"] for k in kids), 2),
            "total_hours": round(total, 2), "billable_hours": round(billable, 2),
            "non_billable_hours": round(total - billable, 2),
            "planned_fact_hours": round(sum(k["planned_fact_hours"] for k in kids), 2),
            "children": kids,
        })
    projects.sort(key=lambda p: (-p["plan_hours"], -p["total_hours"], p["name"].lower()))

    # --- второй разрез: сотрудник → проект → задача ---
    # План сотрудника — оценки задач, где он ответственный. Берутся только
    # «листья» плана (без оценок глубже), иначе этап и подзадачи одного
    # ответственного сложились бы дважды.
    per_employee: Dict[str, Dict[str, List[float]]] = {}
    for task_id, node in nodes.items():
        for emp_id, v in node["emps"].items():
            per_employee.setdefault(emp_id, {})[task_id] = [v[0], v[1], 0.0]
    for task_id, node in nodes.items():
        if node["own_plan"] and task_id not in planned_parents and node["responsible_id"]:
            emp_id = normalize_employee_id(node["responsible_id"])
            per_employee.setdefault(emp_id, {}).setdefault(task_id, [0.0, 0.0, 0.0])[2] = node["own_plan"]

    def summed(type_: str, id_: str, name: str, items: List[Dict[str, Any]], **extra) -> Dict[str, Any]:
        total = sum(i["total_hours"] for i in items)
        billable = sum(i["billable_hours"] for i in items)
        return {
            "type": type_, "id": id_, "name": name,
            "plan_hours": round(sum(i["plan_hours"] for i in items), 2),
            "total_hours": round(total, 2), "billable_hours": round(billable, 2),
            "non_billable_hours": round(total - billable, 2),
            "planned_fact_hours": round(sum(i["planned_fact_hours"] for i in items), 2),
            "children": items, **extra,
        }

    employees: List[Dict[str, Any]] = []
    for emp_id, task_map in per_employee.items():
        by_project: Dict[str, List[Dict[str, Any]]] = {}
        for task_id, (total, billable, plan) in task_map.items():
            node = nodes[task_id]
            by_project.setdefault(node["project"], []).append({
                "type": "task", "id": task_id, "name": node["title"],
                "plan_hours": round(plan, 2), "own_plan_hours": round(plan, 2), "children_plan_hours": 0.0,
                "total_hours": round(total, 2), "billable_hours": round(billable, 2),
                "non_billable_hours": round(total - billable, 2),
                "planned_fact_hours": round(total if plan else 0.0, 2),
                "responsible_name": "", "is_closed": node["closed"], "children": [], "employees": [],
            })
        project_nodes = []
        for project, items in by_project.items():
            items.sort(key=lambda i: (-i["total_hours"], -i["plan_hours"]))
            project_nodes.append(summed(
                "project", project,
                project_names.get(project) or ("Без проекта" if project == NO_PROJECT_KEY else f"Проект {project}"),
                items,
            ))
        project_nodes.sort(key=lambda p: (-p["total_hours"], -p["plan_hours"]))
        employees.append(summed("employee", emp_id, resolve_employee_name(user_map, emp_id), project_nodes))
    employees.sort(key=lambda e: (-e["total_hours"], -e["plan_hours"], e["name"].lower()))

    return {"projects": projects, "employees": employees}


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------

_HEAD = ("Наименование", "Уровень", "Ответственный", "План, ч", "Факт всего, ч", "Факт учит., ч",
         "Факт неучит., ч", "Остаток по задачам с планом, ч", "Примечание")
_WIDTHS = (70, 12, 24, 10, 13, 13, 14, 16, 44)
_FILL_HEAD = PatternFill("solid", fgColor="DCE6F5")
_FILL_TOP = PatternFill("solid", fgColor="EEF2F4")
_LEVELS = {"project": "Проект", "task": "Задача", "employee": "Сотрудник"}


def _count_rows(nodes: Sequence[Mapping[str, Any]]) -> int:
    return sum(
        1 + len(node.get("employees") or []) + _count_rows(node.get("children") or [])
        for node in nodes
    )


def _write_sheet(ws, title: str, roots: Sequence[Mapping[str, Any]]) -> None:
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(list(_HEAD))
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.fill = _FILL_HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    def write(node: Mapping[str, Any], depth: int, level: str, note: str = "") -> None:
        plan = float(node.get("plan_hours") or 0)
        total = float(node.get("total_hours") or 0)
        planned_fact = float(node.get("planned_fact_hours") if node.get("planned_fact_hours") is not None else total)
        name = str(node.get("name") or "")
        # Значение, начинающееся с = + - @, Excel сочтёт формулой.
        if name[:1] in ("=", "+", "-", "@"):
            name = "'" + name
        ws.append([
            name, level, node.get("responsible_name") or None, plan or None, total,
            float(node.get("billable_hours") or 0), float(node.get("non_billable_hours") or 0),
            round(plan - planned_fact, 2) if plan else None, note or None,
        ])
        row = ws.max_row
        ws.cell(row, 1).alignment = Alignment(indent=min(depth * 2, 30), wrap_text=True)
        ws.row_dimensions[row].outline_level = min(depth, 7)
        if depth == 0:
            for cell in ws[row]:
                cell.font = Font(bold=True)
                cell.fill = _FILL_TOP
        if plan and planned_fact > plan:
            ws.cell(row, 8).font = Font(bold=True, color="C0392B")

    def walk(node: Mapping[str, Any], depth: int) -> None:
        notes = []
        own, kids = float(node.get("own_plan_hours") or 0), float(node.get("children_plan_hours") or 0)
        if own and kids and abs(own - kids) > 0.01:
            notes.append(f"оценка этапа {own:g} ≠ сумма подзадач {kids:g}")
        if node.get("is_closed"):
            notes.append("завершена")
        write(node, depth, _LEVELS.get(str(node.get("type")), ""), "; ".join(notes))
        for employee in node.get("employees") or []:
            write(employee, depth + 1, _LEVELS["employee"])
        for child in node.get("children") or []:
            walk(child, depth + 1)

    total_row = {
        "name": "ИТОГО",
        "plan_hours": sum(r["plan_hours"] for r in roots),
        "total_hours": sum(r["total_hours"] for r in roots),
        "billable_hours": sum(r["billable_hours"] for r in roots),
        "non_billable_hours": sum(r["non_billable_hours"] for r in roots),
        "planned_fact_hours": sum(r["planned_fact_hours"] for r in roots),
    }
    write(total_row, 0, "Итог", "остаток — только по задачам с планом")
    for root in roots:
        walk(root, 0)

    for index, width in enumerate(_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "B3"
    ws.sheet_properties.outlinePr.summaryBelow = False
    for row in ws.iter_rows(min_row=3, min_col=4, max_col=8):
        for cell in row:
            cell.number_format = "#,##0.0"


def build_plan_fact_workbook(report: Mapping[str, Any], *, date_from: str = "", date_to: str = "") -> io.BytesIO:
    projects = report.get("projects") or []
    employees = report.get("employees") or []
    if _count_rows(projects) + _count_rows(employees) > MAX_EXPORT_ROWS:
        raise ExportTooLargeError(_count_rows(projects) + _count_rows(employees))

    period = f"{date_from} — {date_to}".strip(" —")
    suffix = f" · факт за период {period}" if period else ""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Проекты-задачи-сотрудники"
    _write_sheet(ws, f"План / факт по проектам, задачам и сотрудникам{suffix}", projects)
    _write_sheet(
        wb.create_sheet("Сотрудники"),
        f"План / факт по сотрудникам{suffix}. План — по задачам, где сотрудник ответственный",
        employees,
    )
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
